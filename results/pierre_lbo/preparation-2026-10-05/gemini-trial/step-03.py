import openpyxl

wb = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=True)

print("=== FORECAST VALUES ===")
ws = wb['Forecast']
for r in range(4, 13):
    vals = [ws.cell(r, c).value for c in range(1, 21)]
    print(f"Row {r:2d}: {vals}")

print("\n=== DEBT VALUES ===")
ws = wb['Debt']
for r in range(4, 13):
    vals = [ws.cell(r, c).value for c in range(1, 29)]
    print(f"Row {r:2d}: {vals}")

print("\n=== RETURNS VALUES ===")
ws = wb['Returns']
for r in range(4, 40):
    k = ws.cell(r, 1).value
    v = ws.cell(r, 2).value
    if k is not None:
        print(f"{k}: {v}")
