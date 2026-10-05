import json
import openpyxl

with open('inputs/base_case.json') as f:
    base_case = json.load(f)

print("Base case keys:", list(base_case.keys()))

wb = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=False)
print("Sheet names:", wb.sheetnames)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"\n--- Sheet: {sheetname} ({ws.max_row}x{ws.max_column}) ---")
    for r in range(1, min(ws.max_row + 1, 35)):
        row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 12))]
        if any(v is not None for v in row_vals):
            print(f"Row {r:2d}: {row_vals}")
