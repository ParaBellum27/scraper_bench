import openpyxl

wb = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=False)
ws = wb['Debt']

headers = [ws.cell(4, c).value for c in range(1, 29)]
row5 = [ws.cell(5, c).value for c in range(1, 29)]

for c, (h, f) in enumerate(zip(headers, row5), 1):
    col_letter = openpyxl.utils.get_column_letter(c)
    print(f"Col {col_letter} ({c:2d}): {h} => {f}")
