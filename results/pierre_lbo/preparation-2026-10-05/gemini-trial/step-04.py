import openpyxl

wb = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=True)
ws_debt = wb['Debt']

# Check how many rows are in Debt
print("Debt row 5 to 12 row status:")
for r in range(5, 13):
    period = ws_debt.cell(r, 1).value
    eval_row = ws_debt.cell(r, 2).value
    status = ws_debt.cell(r, 18).value
    tlb_bal = ws_debt.cell(r, 26).value
    print(f"Row {r}: period={period}, eval_row={eval_row}, status={status}, tlb_bal={tlb_bal}")
