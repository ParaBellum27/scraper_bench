import openpyxl

wb = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=False)
ws = wb['Debt']
for r in range(1, 13):
    row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column+1)]
    while row_vals and row_vals[-1] is None:
        row_vals.pop()
    print(f"Row {r:2d}: {row_vals}")
