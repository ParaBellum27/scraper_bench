import json, sys, calendar
from datetime import date, datetime
import openpyxl

wb = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=True)
forecast = wb['Forecast']
debt = wb['Debt']
returns = wb['Returns']
inputs = wb['Inputs']

exit_p = inputs['B8'].value

with open('inputs/base_case.json') as f:
    inp = json.load(f)

# Import solve from solution.py
import solution
res = solution.solve(inp)

# Expected values from workbook:
expected = {}

expected['status'] = returns['B16'].value
expected['failure_period'] = returns['B17'].value
expected['funding_shortfall'] = returns['B18'].value
expected['dates'] = [forecast.cell(4+p, 3).value.strftime("%Y-%m-%d") for p in range(1, exit_p+1)]
expected['operating_fractions'] = [forecast.cell(4+p, 5).value for p in range(1, exit_p+1)]
expected['interest_year_fractions'] = [forecast.cell(4+p, 6).value for p in range(1, exit_p+1)]
expected['exit_date'] = returns['B13'].value.strftime("%Y-%m-%d")
expected['holding_period_years'] = returns['B14'].value

expected['entry_ebitda'] = returns['B5'].value
expected['entry_ev'] = returns['B6'].value
expected['entry_equity'] = returns['B7'].value
expected['ordinary_equity'] = returns['B8'].value
expected['preferred_equity'] = returns['B9'].value
expected['sponsor_investment'] = returns['B10'].value
expected['management_investment'] = returns['B11'].value

expected['revenue'] = [forecast.cell(4+p, 12).value for p in range(1, exit_p+1)]
expected['gross_profit'] = [forecast.cell(4+p, 13).value for p in range(1, exit_p+1)]
expected['opex'] = [forecast.cell(4+p, 14).value for p in range(1, exit_p+1)]
expected['ebitda'] = [forecast.cell(4+p, 15).value for p in range(1, exit_p+1)]
expected['ebitda_margin'] = [forecast.cell(4+p, 16).value for p in range(1, exit_p+1)]
expected['working_capital'] = [forecast.cell(4+p, 17).value for p in range(1, exit_p+1)]
expected['change_working_capital'] = [forecast.cell(4+p, 18).value for p in range(1, exit_p+1)]
expected['cash_capex'] = [forecast.cell(4+p, 19).value for p in range(1, exit_p+1)]
expected['cash_ebitda'] = [forecast.cell(4+p, 20).value for p in range(1, exit_p+1)]

expected['cash_interest'] = [debt.cell(4+p, 8).value for p in range(1, exit_p+1)]
expected['tlb_interest'] = [debt.cell(4+p, 6).value for p in range(1, exit_p+1)]
expected['rcf_interest'] = [debt.cell(4+p, 7).value for p in range(1, exit_p+1)]
expected['cash_taxes'] = [debt.cell(4+p, 10).value for p in range(1, exit_p+1)]
expected['pre_financing_cash_flow'] = [debt.cell(4+p, 11).value for p in range(1, exit_p+1)]
expected['tlb_pik'] = [debt.cell(4+p, 9).value for p in range(1, exit_p+1)]
expected['mandatory_tlb_repayment'] = [debt.cell(4+p, 12).value for p in range(1, exit_p+1)]
expected['mandatory_rcf_repayment'] = [debt.cell(4+p, 13).value for p in range(1, exit_p+1)]
expected['rcf_draw'] = [debt.cell(4+p, 19).value for p in range(1, exit_p+1)]
expected['rcf_repayment'] = [debt.cell(4+p, 23).value for p in range(1, exit_p+1)]
expected['tlb_sweep'] = [debt.cell(4+p, 25).value for p in range(1, exit_p+1)]
expected['tlb_balance'] = [debt.cell(4+p, 26).value for p in range(1, exit_p+1)]
expected['rcf_balance'] = [debt.cell(4+p, 27).value for p in range(1, exit_p+1)]
expected['cash_balance'] = [debt.cell(4+p, 28).value for p in range(1, exit_p+1)]

expected['exit_ev'] = returns['B20'].value
expected['exit_fees'] = returns['B21'].value
expected['exit_debt'] = returns['B22'].value
expected['exit_cash'] = returns['B23'].value
expected['equity_value_before_floor'] = returns['B24'].value
expected['equity_proceeds'] = returns['B25'].value

expected['preferred_accrued'] = returns['B27'].value
expected['preferred_payout'] = returns['B28'].value
expected['ordinary_payout'] = returns['B29'].value
expected['sponsor_proceeds'] = returns['B30'].value
expected['management_proceeds'] = returns['B31'].value

expected['deal_mom'] = returns['B33'].value
expected['sponsor_mom'] = returns['B34'].value
expected['management_mom'] = returns['B35'].value
expected['deal_irr'] = returns['B37'].value
expected['sponsor_irr'] = returns['B38'].value
expected['management_irr'] = returns['B39'].value

diffs = []
for k in res:
    if k not in expected:
        diffs.append(f"Key {k} not in expected!")
        continue
    exp_v = expected[k]
    act_v = res[k]
    
    if isinstance(act_v, list):
        if len(act_v) != len(exp_v):
            diffs.append(f"{k}: len {len(act_v)} != {len(exp_v)}")
        else:
            for i, (a, b) in enumerate(zip(act_v, exp_v)):
                if isinstance(a, str):
                    if a != b:
                        diffs.append(f"{k}[{i}]: {a} != {b}")
                else:
                    if abs(a - b) > 1e-7 * max(abs(b), 1.0):
                        diffs.append(f"{k}[{i}]: {a} != {b} (diff {abs(a-b)})")
    else:
        if exp_v is None or act_v is None:
            if exp_v != act_v:
                diffs.append(f"{k}: {act_v} != {exp_v}")
        elif isinstance(act_v, str):
            if act_v != exp_v:
                diffs.append(f"{k}: {act_v} != {exp_v}")
        else:
            if abs(act_v - exp_v) > 1e-7 * max(abs(exp_v), 1.0):
                diffs.append(f"{k}: {act_v} != {exp_v} (diff {abs(act_v - exp_v)})")

print(f"Total checked fields: {len(expected)}")
if not diffs:
    print("ALL 55 FIELDS MATCH PERFECTLY!")
else:
    print(f"Found {len(diffs)} differences:")
    for d in diffs:
        print(" ", d)
