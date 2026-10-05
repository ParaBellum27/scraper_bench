import openpyxl

wb = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=True)

forecast = wb['Forecast']
debt = wb['Debt']
returns = wb['Returns']
inputs = wb['Inputs']

exit_p = inputs['B8'].value
print(f"exit_period = {exit_p}")

# Let's see what rows are populated for Forecast and Debt
print("\nForecast columns for p=1..exit_p:")
for col_idx, col_name in [
    (3, 'dates'), (5, 'operating_fractions'), (6, 'interest_year_fractions'),
    (12, 'revenue'), (13, 'gross_profit'), (14, 'opex'), (15, 'ebitda'),
    (16, 'ebitda_margin'), (17, 'working_capital'), (18, 'change_working_capital'),
    (19, 'cash_capex'), (20, 'cash_ebitda')
]:
    vals = [forecast.cell(4+p, col_idx).value for p in range(1, exit_p+1)]
    print(f"  {col_name}: {vals}")

print("\nDebt columns for p=1..exit_p:")
# Financing arrays:
# cash_interest, tlb_interest, rcf_interest, cash_taxes, pre_financing_cash_flow, tlb_pik,
# mandatory_tlb_repayment, mandatory_rcf_repayment, rcf_draw, rcf_repayment, tlb_sweep,
# tlb_balance, rcf_balance, cash_balance
debt_map = {
    'tlb_interest': 6, # Col F
    'rcf_interest': 7, # Col G
    'cash_interest': 8, # Col H
    'tlb_pik': 9, # Col I
    'cash_taxes': 10, # Col J
    'pre_financing_cash_flow': 11, # Col K
    'mandatory_tlb_repayment': 12, # Col L
    'mandatory_rcf_repayment': 13, # Col M
    'rcf_draw': 19, # Col S
    'rcf_repayment': 23, # Col W
    'tlb_sweep': 25, # Col Y
    'tlb_balance': 26, # Col Z
    'rcf_balance': 27, # Col AA
    'cash_balance': 28, # Col AB
}
for name, col_idx in debt_map.items():
    vals = [debt.cell(4+p, col_idx).value for p in range(1, exit_p+1)]
    print(f"  {name}: {vals}")

print("\nReturns sheet values:")
for r in range(5, 40):
    k = returns.cell(r, 1).value
    v = returns.cell(r, 2).value
    if k:
        print(f"  {k}: {v}")
