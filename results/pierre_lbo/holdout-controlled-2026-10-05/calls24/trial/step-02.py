import openpyxl

wb = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=False)
wb_val = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=True)

for name in wb.sheetnames:
    ws = wb[name]
    ws_val = wb_val[name]
    print(f"\n--- SHEET: {name} (max_row={ws.max_row}, max_column={ws.max_column}) ---")
    for r in range(1, min(ws.max_row+1, 35)):
        row_vals = []
        for c in range(1, min(ws.max_column+1, 15)):
            f = ws.cell(r, c).value
            v = ws_val.cell(r, c).value
            if f is not None or v is not None:
                if str(f).startswith('='):
                    row_vals.append(f"C{c}: {f} => {v}")
                else:
                    row_vals.append(f"C{c}: {v}")
        if row_vals:
            print(f"R{r}: " + " | ".join(row_vals[:8]))
