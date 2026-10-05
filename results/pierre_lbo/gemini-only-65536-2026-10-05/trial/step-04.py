import openpyxl

wb_f = openpyxl.load_workbook("pierre_lbo_one.xlsx", data_only=False)
ws = wb_f["Returns"]

for r in range(1, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if any(vals):
        print(f"Row {r:2d}: {vals}")
