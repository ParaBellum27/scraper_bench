import openpyxl

wb_f = openpyxl.load_workbook("pierre_lbo_one.xlsx", data_only=False)

def dump_sheet(ws, max_cols=35):
    print(f"=== Sheet: {ws.title} ===")
    headers = [ws.cell(4, c).value for c in range(1, max_cols + 1)]
    print("Headers:", [(i+1, h) for i, h in enumerate(headers) if h is not None])
    for r in range(5, 7):
        print(f"Row {r}:")
        for c in range(1, max_cols + 1):
            val = ws.cell(r, c).value
            h = headers[c-1]
            if val is not None:
                print(f"  Col {openpyxl.utils.get_column_letter(c)} ({h}): {val}")

dump_sheet(wb_f["Forecast"])
dump_sheet(wb_f["Debt"])
