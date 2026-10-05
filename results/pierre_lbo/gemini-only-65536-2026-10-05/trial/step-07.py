import openpyxl, json

wb = openpyxl.load_workbook("pierre_lbo_one.xlsx", data_only=True)
ws_ret = wb["Returns"]
ws_fc = wb["Forecast"]
ws_debt = wb["Debt"]

print("--- RETURNS SHEET ---")
for r in range(5, 40):
    k = ws_ret.cell(r, 1).value
    v = ws_ret.cell(r, 2).value
    if k:
        print(f"  {k}: {v}")

print("\n--- FORECAST (first 5) ---")
for col_idx, name in [(12, 'revenue'), (13, 'gross_profit'), (14, 'opex'), (15, 'ebitda'), (16, 'ebitda_margin'), (17, 'working_capital'), (18, 'change_working_capital'), (19, 'cash_capex'), (20, 'cash_ebitda')]:
    vals = [ws_fc.cell(r, col_idx).value for r in range(5, 10)]
    print(f"  {name}: {vals}")

print("\n--- TIMING (first 5) ---")
for col_idx, name in [(3, 'dates'), (5, 'operating_fractions'), (6, 'interest_year_fractions')]:
    vals = [ws_fc.cell(r, col_idx).value for r in range(5, 10)]
    print(f"  {name}: {vals}")

print("\n--- DEBT (first 5) ---")
debt_cols = [
    (8, 'cash_interest'), (6, 'tlb_interest'), (7, 'rcf_interest'), (10, 'cash_taxes'),
    (11, 'pre_financing_cash_flow'), (9, 'tlb_pik'), (12, 'mandatory_tlb_repayment'),
    (13, 'mandatory_rcf_repayment'), (19, 'rcf_draw'), (23, 'rcf_repayment'),
    (25, 'tlb_sweep'), (26, 'tlb_balance'), (27, 'rcf_balance'), (28, 'cash_balance')
]
for col_idx, name in debt_cols:
    vals = [ws_debt.cell(r, col_idx).value for r in range(5, 10)]
    print(f"  {name}: {vals}")
