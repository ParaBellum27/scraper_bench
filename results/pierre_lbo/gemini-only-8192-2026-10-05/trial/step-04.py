import openpyxl

wb = openpyxl.load_workbook("pierre_lbo_one.xlsx", data_only=True)
ws_forecast = wb["Forecast"]
for r in range(4, 13):
    row_vals = [ws_forecast.cell(r, c).value for c in range(1, 21)]
    print(f"R{r}: {row_vals}")

ws_debt = wb["Debt"]
for r in range(4, 13):
    row_vals = [ws_debt.cell(r, c).value for c in range(1, 29)]
    print(f"Debt R{r}: {row_vals[:12]}")
    print(f"        : {row_vals[12:]}")

ws_returns = wb["Returns"]
for r in range(4, 46):
    print(f"Ret R{r}: {ws_returns.cell(r, 1).value} = {ws_returns.cell(r, 2).value}")
