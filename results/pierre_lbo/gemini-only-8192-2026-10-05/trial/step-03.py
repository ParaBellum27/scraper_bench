import openpyxl

wb = openpyxl.load_workbook("pierre_lbo_one.xlsx", data_only=False)

def dump_sheet(sheetname):
    ws = wb[sheetname]
    print(f"=== {sheetname} ===")
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        while vals and vals[-1] is None:
            vals.pop()
        if vals:
            print(f"R{r:02d}: {vals}")

dump_sheet("Forecast")
dump_sheet("Debt")
dump_sheet("Returns")
