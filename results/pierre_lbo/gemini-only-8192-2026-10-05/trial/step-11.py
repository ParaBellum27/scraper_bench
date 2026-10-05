import json
import datetime
import calendar
import openpyxl

def solve(data):
    case = data["case"]
    last_fiscal_year_end = datetime.date.fromisoformat(data["last_fiscal_year_end"])
    deal_date = datetime.date.fromisoformat(data["deal_date"])
    exit_period = data["exit_period"]
    opening_revenue = float(data["opening_revenue"])
    opening_working_capital = float(data["opening_working_capital"])
    entry_multiple = float(data["entry_multiple"])
    exit_multiple = float(data["exit_multiple"])
    exit_fee_rate = float(data["exit_fee_rate"])
    tax_rate = float(data["tax_rate"])
    ordinary_equity_fraction = float(data["ordinary_equity_fraction"])
    management_ordinary_fraction = float(data["management_ordinary_fraction"])
    preferred_pik_rate = float(data["preferred_pik_rate"])
    working_capital_ratio = [float(x) for x in data["working_capital_ratio"]]

    term_loan = data["term_loan"]
    tlb_principal = float(term_loan["principal"])
    tlb_cash_rate = float(term_loan["cash_rate"])
    tlb_pik_rate = float(term_loan["pik_rate"])
    tlb_amort_rate = float(term_loan["amortization_rate"])
    tlb_cash_sweep = float(term_loan["cash_sweep"])
    tlb_maturity = int(term_loan["maturity_period"])

    revolver = data["revolver"]
    rcf_commitment = float(revolver["commitment"])
    rcf_cash_rate = float(revolver["cash_rate"])
    rcf_maturity = int(revolver["maturity_period"])

    scen = data["scenarios"][case - 1]
    rev_growth = [float(x) for x in scen["revenue_growth"]]
    gp_margin = [float(x) for x in scen["gross_profit_margin"]]
    opex_margin = [float(x) for x in scen["opex_margin"]]
    capex = [float(x) for x in scen["capex"]]

    def advance_edate_12(d):
        y = d.year + 1
        m = d.month
        max_d = calendar.monthrange(y, m)[1]
        return datetime.date(y, m, min(d.day, max_d))

    dates_all = []
    prev_d = last_fiscal_year_end
    for _ in range(8):
        next_d = advance_edate_12(prev_d)
        dates_all.append(next_d)
        prev_d = next_d

    operating_fractions_all = []
    interest_year_fractions_all = []
    prev_d = last_fiscal_year_end
    for i in range(8):
        c_date = dates_all[i]
        d_date = max(prev_d, deal_date)
        op_frac = (c_date - d_date).days / (c_date - prev_d).days
        int_frac = (c_date - d_date).days / 365.0
        operating_fractions_all.append(op_frac)
        interest_year_fractions_all.append(int_frac)
        prev_d = c_date

    revenue_all = []
    gross_profit_all = []
    opex_all = []
    ebitda_all = []
    ebitda_margin_all = []
    working_capital_all = []
    change_working_capital_all = []
    cash_capex_all = []
    cash_ebitda_all = []

    curr_rev = opening_revenue
    curr_wc = opening_working_capital
    for i in range(8):
        curr_rev = curr_rev * (1.0 + rev_growth[i])
        revenue_all.append(curr_rev)
        gp = curr_rev * gp_margin[i]
        gross_profit_all.append(gp)
        op = curr_rev * opex_margin[i]
        opex_all.append(op)
        eb = gp - op
        ebitda_all.append(eb)
        ebitda_margin_all.append(gp_margin[i] - opex_margin[i])
        wc = curr_rev * working_capital_ratio[i]
        working_capital_all.append(wc)
        ch_wc = wc - curr_wc
        change_working_capital_all.append(ch_wc)
        curr_wc = wc
        cash_capex_all.append(capex[i] * operating_fractions_all[i])
        cash_ebitda_all.append(eb * operating_fractions_all[i])

    entry_ebitda = ebitda_all[0]
    entry_ev = entry_ebitda * entry_multiple
    entry_equity = entry_ev - tlb_principal
    ordinary_equity = entry_equity * ordinary_equity_fraction
    preferred_equity = entry_equity - ordinary_equity
    management_investment = ordinary_equity * management_ordinary_fraction
    sponsor_investment = entry_equity - management_investment

    exit_date_str = dates_all[exit_period - 1].isoformat()
    holding_period_years = (dates_all[exit_period - 1] - deal_date).days / 365.0

    # Financing loop
    tlb_interest_list = []
    rcf_interest_list = []
    cash_interest_list = []
    tlb_pik_list = []
    cash_taxes_list = []
    pre_financing_cash_flow_list = []
    mandatory_tlb_list = []
    mandatory_rcf_list = []
    rcf_draw_list = []
    rcf_repayment_list = []
    tlb_sweep_list = []
    tlb_balance_list = []
    rcf_balance_list = []
    cash_balance_list = []

    cur_tlb = tlb_principal
    cur_rcf = 0.0
    cur_cash = 0.0

    status = "ok"
    failure_period = None
    funding_shortfall = 0.0

    for p in range(1, exit_period + 1):
        idx = p - 1
        tlb_int = cur_tlb * tlb_cash_rate * interest_year_fractions_all[idx]
        rcf_int = cur_rcf * rcf_cash_rate * interest_year_fractions_all[idx]
        c_int = tlb_int + rcf_int
        tlb_p = cur_tlb * tlb_pik_rate * interest_year_fractions_all[idx]
        c_tax = max(cash_ebitda_all[idx] - c_int, 0.0) * tax_rate
        pre_cf = cash_ebitda_all[idx] - cash_capex_all[idx] - change_working_capital_all[idx] - c_tax

        # Candidate repayments
        if p >= tlb_maturity:
            cand_mand_tlb = cur_tlb + tlb_p
        else:
            cand_mand_tlb = min(cur_tlb + tlb_p, tlb_principal * tlb_amort_rate * operating_fractions_all[idx])

        if p >= rcf_maturity:
            cand_mand_rcf = cur_rcf
        else:
            cand_mand_rcf = 0.0

        cash_before_rep = cur_cash + pre_cf - c_int
        req_draw = max(cand_mand_tlb + cand_mand_rcf - cash_before_rep, 0.0)

        if p < rcf_maturity:
            avail_draw = max(rcf_commitment - cur_rcf, 0.0)
        else:
            avail_draw = 0.0

        shortfall = max(req_draw - avail_draw, 0.0)
        liq_scale = max(
            cur_cash,
            abs(cash_ebitda_all[idx]),
            cash_capex_all[idx],
            abs(change_working_capital_all[idx]),
            c_tax,
            c_int,
            cand_mand_tlb + cand_mand_rcf,
            avail_draw
        )

        if shortfall > 1e-12 * liq_scale:
            status = "liquidity_shortfall"
            failure_period = p
            funding_shortfall = shortfall
            break

        actual_draw = min(req_draw, avail_draw)
        cash_after_mand = max(cash_before_rep + actual_draw - cand_mand_tlb - cand_mand_rcf, 0.0)
        interim_rcf = cur_rcf + actual_draw - cand_mand_rcf
        vol_rcf = min(cash_after_mand, interim_rcf)
        total_rcf_rep = cand_mand_rcf + vol_rcf
        cash_after_rcf = cash_after_mand - vol_rcf
        sweep = min(cash_after_rcf * tlb_cash_sweep, cur_tlb + tlb_p - cand_mand_tlb)
        new_tlb = cur_tlb + tlb_p - cand_mand_tlb - sweep
        new_rcf = interim_rcf - vol_rcf
        new_cash = cash_after_rcf - sweep

        tlb_interest_list.append(tlb_int)
        rcf_interest_list.append(rcf_int)
        cash_interest_list.append(c_int)
        tlb_pik_list.append(tlb_p)
        cash_taxes_list.append(c_tax)
        pre_financing_cash_flow_list.append(pre_cf)
        mandatory_tlb_list.append(cand_mand_tlb)
        mandatory_rcf_list.append(cand_mand_rcf)
        rcf_draw_list.append(actual_draw)
        rcf_repayment_list.append(total_rcf_rep)
        tlb_sweep_list.append(sweep)
        tlb_balance_list.append(new_tlb)
        rcf_balance_list.append(new_rcf)
        cash_balance_list.append(new_cash)

        cur_tlb = new_tlb
        cur_rcf = new_rcf
        cur_cash = new_cash

    out_dates = [d.isoformat() for d in dates_all[:exit_period]]
    out_op_frac = operating_fractions_all[:exit_period]
    out_int_frac = interest_year_fractions_all[:exit_period]
    out_rev = revenue_all[:exit_period]
    out_gp = gross_profit_all[:exit_period]
    out_op = opex_all[:exit_period]
    out_eb = ebitda_all[:exit_period]
    out_eb_margin = ebitda_margin_all[:exit_period]
    out_wc = working_capital_all[:exit_period]
    out_ch_wc = change_working_capital_all[:exit_period]
    out_cash_capex = cash_capex_all[:exit_period]
    out_cash_eb = cash_ebitda_all[:exit_period]

    if status == "ok":
        exit_ev = ebitda_all[exit_period - 1] * exit_multiple
        exit_fees = max(exit_ev, 0.0) * exit_fee_rate
        exit_debt = cur_tlb + cur_rcf
        exit_cash = cur_cash
        equity_value_before_floor = exit_ev - exit_fees - exit_debt + exit_cash
        equity_proceeds = max(equity_value_before_floor, 0.0)

        pref_acc = preferred_equity * ((1.0 + preferred_pik_rate) ** holding_period_years)
        if equity_proceeds - pref_acc <= 1e-12 * equity_proceeds:
            pref_payout = equity_proceeds
        else:
            pref_payout = pref_acc
        ord_payout = equity_proceeds - pref_payout
        spon_proceeds = pref_payout + ord_payout * (1.0 - management_ordinary_fraction)
        mgmt_proceeds = ord_payout * management_ordinary_fraction

        deal_mom = (equity_proceeds / entry_equity) if entry_equity > 0 else None
        sponsor_mom = (spon_proceeds / sponsor_investment) if sponsor_investment > 0 else None
        mgmt_mom = (mgmt_proceeds / management_investment) if management_investment > 0 else None

        deal_irr = ((equity_proceeds / entry_equity) ** (1.0 / holding_period_years) - 1.0) if (entry_equity > 0 and equity_proceeds > 0) else None
        sponsor_irr = ((spon_proceeds / sponsor_investment) ** (1.0 / holding_period_years) - 1.0) if (sponsor_investment > 0 and spon_proceeds > 0) else None
        mgmt_irr = ((mgmt_proceeds / management_investment) ** (1.0 / holding_period_years) - 1.0) if (management_investment > 0 and mgmt_proceeds > 0) else None
    else:
        exit_ev = None
        exit_fees = None
        exit_debt = None
        exit_cash = None
        equity_value_before_floor = None
        equity_proceeds = None
        pref_acc = None
        pref_payout = None
        ord_payout = None
        spon_proceeds = None
        mgmt_proceeds = None
        deal_mom = None
        sponsor_mom = None
        mgmt_mom = None
        deal_irr = None
        sponsor_irr = None
        mgmt_irr = None

    return {
        "status": status,
        "failure_period": failure_period,
        "funding_shortfall": funding_shortfall,
        "dates": out_dates,
        "operating_fractions": out_op_frac,
        "interest_year_fractions": out_int_frac,
        "exit_date": exit_date_str,
        "holding_period_years": holding_period_years,

        "entry_ebitda": entry_ebitda,
        "entry_ev": entry_ev,
        "entry_equity": entry_equity,
        "ordinary_equity": ordinary_equity,
        "preferred_equity": preferred_equity,
        "sponsor_investment": sponsor_investment,
        "management_investment": management_investment,

        "revenue": out_rev,
        "gross_profit": out_gp,
        "opex": out_op,
        "ebitda": out_eb,
        "ebitda_margin": out_eb_margin,
        "working_capital": out_wc,
        "change_working_capital": out_ch_wc,
        "cash_capex": out_cash_capex,
        "cash_ebitda": out_cash_eb,

        "cash_interest": cash_interest_list,
        "tlb_interest": tlb_interest_list,
        "rcf_interest": rcf_interest_list,
        "cash_taxes": cash_taxes_list,
        "pre_financing_cash_flow": pre_financing_cash_flow_list,
        "tlb_pik": tlb_pik_list,
        "mandatory_tlb_repayment": mandatory_tlb_list,
        "mandatory_rcf_repayment": mandatory_rcf_list,
        "rcf_draw": rcf_draw_list,
        "rcf_repayment": rcf_repayment_list,
        "tlb_sweep": tlb_sweep_list,
        "tlb_balance": tlb_balance_list,
        "rcf_balance": rcf_balance_list,
        "cash_balance": cash_balance_list,

        "exit_ev": exit_ev,
        "exit_fees": exit_fees,
        "exit_debt": exit_debt,
        "exit_cash": exit_cash,
        "equity_value_before_floor": equity_value_before_floor,
        "equity_proceeds": equity_proceeds,

        "preferred_accrued": pref_acc,
        "preferred_payout": pref_payout,
        "ordinary_payout": ord_payout,
        "sponsor_proceeds": spon_proceeds,
        "management_proceeds": mgmt_proceeds,
        "deal_mom": deal_mom,
        "sponsor_mom": sponsor_mom,
        "management_mom": mgmt_mom,
        "deal_irr": deal_irr,
        "sponsor_irr": sponsor_irr,
        "management_irr": mgmt_irr,
    }

with open("inputs/base_case.json") as f:
    inp = json.load(f)

res = solve(inp)
print("Keys count:", len(res))

# Now compare with workbook
wb = openpyxl.load_workbook("pierre_lbo_one.xlsx", data_only=True)
ws_ret = wb["Returns"]
ws_fc = wb["Forecast"]
ws_debt = wb["Debt"]

mismatches = []
def check_val(name, actual, expected):
    if actual is None and expected is None:
        return
    if isinstance(actual, float) and isinstance(expected, (int, float)):
        diff = abs(actual - expected)
        reldiff = diff / (abs(expected) + 1e-12)
        if diff > 1e-6 and reldiff > 1e-6:
            mismatches.append(f"{name}: actual {actual} != expected {expected} (diff {diff})")
    elif actual != expected:
        mismatches.append(f"{name}: actual {actual} != expected {expected}")

# Check scalars from Returns
scalar_map = {
    "entry_ebitda": ws_ret["B5"].value,
    "entry_ev": ws_ret["B6"].value,
    "entry_equity": ws_ret["B7"].value,
    "ordinary_equity": ws_ret["B8"].value,
    "preferred_equity": ws_ret["B9"].value,
    "sponsor_investment": ws_ret["B10"].value,
    "management_investment": ws_ret["B11"].value,
    "exit_date": ws_ret["B13"].value.strftime("%Y-%m-%d"),
    "holding_period_years": ws_ret["B14"].value,
    "status": ws_ret["B16"].value,
    "failure_period": ws_ret["B17"].value if ws_ret["B17"].value != "" else None,
    "funding_shortfall": ws_ret["B18"].value,
    "exit_ev": ws_ret["B20"].value,
    "exit_fees": ws_ret["B21"].value,
    "exit_debt": ws_ret["B22"].value,
    "exit_cash": ws_ret["B23"].value,
    "equity_value_before_floor": ws_ret["B24"].value,
    "equity_proceeds": ws_ret["B25"].value,
    "preferred_accrued": ws_ret["B27"].value,
    "preferred_payout": ws_ret["B28"].value,
    "ordinary_payout": ws_ret["B29"].value,
    "sponsor_proceeds": ws_ret["B30"].value,
    "management_proceeds": ws_ret["B31"].value,
    "deal_mom": ws_ret["B33"].value,
    "sponsor_mom": ws_ret["B34"].value,
    "management_mom": ws_ret["B35"].value,
    "deal_irr": ws_ret["B37"].value,
    "sponsor_irr": ws_ret["B38"].value,
    "management_irr": ws_ret["B39"].value,
}

for k, exp in scalar_map.items():
    check_val(k, res[k], exp)

# Check forecast arrays
fc_cols = {
    "dates": (3, lambda v: v.strftime("%Y-%m-%d")),
    "operating_fractions": (5, float),
    "interest_year_fractions": (6, float),
    "revenue": (12, float),
    "gross_profit": (13, float),
    "opex": (14, float),
    "ebitda": (15, float),
    "ebitda_margin": (16, float),
    "working_capital": (17, float),
    "change_working_capital": (18, float),
    "cash_capex": (19, float),
    "cash_ebitda": (20, float),
}

for k, (col, conv) in fc_cols.items():
    for p in range(5):
        val = conv(ws_fc.cell(5 + p, col).value)
        check_val(f"{k}[{p}]", res[k][p], val)

debt_cols = {
    "tlb_interest": 6,
    "rcf_interest": 7,
    "cash_interest": 8,
    "tlb_pik": 9,
    "cash_taxes": 10,
    "pre_financing_cash_flow": 11,
    "mandatory_tlb_repayment": 12,
    "mandatory_rcf_repayment": 13,
    "rcf_draw": 19,
    "rcf_repayment": 23,
    "tlb_sweep": 25,
    "tlb_balance": 26,
    "rcf_balance": 27,
    "cash_balance": 28,
}

for k, col in debt_cols.items():
    for p in range(5):
        val = float(ws_debt.cell(5 + p, col).value)
        check_val(f"{k}[{p}]", res[k][p], val)

print(f"Total mismatches: {len(mismatches)}")
for m in mismatches[:10]:
    print(m)
