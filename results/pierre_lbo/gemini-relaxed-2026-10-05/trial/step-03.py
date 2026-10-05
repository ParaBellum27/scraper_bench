import openpyxl

wb_f = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=False)
wb_v = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=True)

print("=== FORECAST COLUMNS ===")
ws_f = wb_f['Forecast']
ws_v = wb_v['Forecast']
headers_f = [ws_f.cell(4, c).value for c in range(1, ws_f.max_column + 1)]
print("Headers:", list(enumerate(headers_f, 1)))

for r in range(5, 13):
    row_v = [ws_v.cell(r, c).value for c in range(1, ws_f.max_column + 1)]
    row_f = [ws_f.cell(r, c).value for c in range(1, ws_f.max_column + 1)]
    print(f"Row {r}:")
    for c_idx, (h, f, v) in enumerate(zip(headers_f, row_f, row_v), 1):
        if f != v:
            print(f"  Col {c_idx} ({h}): {f}  ==>  {v}")
        else:
            print(f"  Col {c_idx} ({h}): {v}")
