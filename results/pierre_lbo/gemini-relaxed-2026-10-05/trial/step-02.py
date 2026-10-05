import openpyxl

wb_formulas = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=False)
wb_values = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=True)

for sheetname in wb_formulas.sheetnames:
    ws_f = wb_formulas[sheetname]
    ws_v = wb_values[sheetname]
    print(f"=== Sheet: {sheetname} ===")
    for r in range(1, min(ws_f.max_row + 1, 45)):
        row_f = [ws_f.cell(r, c).value for c in range(1, min(ws_f.max_column + 1, 12))]
        row_v = [ws_v.cell(r, c).value for c in range(1, min(ws_v.max_column + 1, 12))]
        if any(x is not None for x in row_f):
            # print non-empty rows
            row_str = " | ".join(f"{str(f)} [{str(v)}]" if f != v else str(f) for f, v in zip(row_f, row_v) if f is not None or v is not None)
            print(f"R{r}: {row_str[:120]}")
