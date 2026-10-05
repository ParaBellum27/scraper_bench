import json
import datetime
import math
import openpyxl

wb = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=True)
ws_forecast = wb['Forecast']
ws_debt = wb['Debt']
ws_returns = wb['Returns']
ws_inputs = wb['Inputs']

exit_p = ws_inputs['B8'].value

# Let's read expected fields from workbook
expected = {}

# Status and timing
expected['status'] = ws_returns['B16'].value
expected['failure_period'] = ws_returns['B17'].value
if expected['failure_period'] == "":
    expected['failure_period'] = None
expected['funding_shortfall'] = ws_returns['B18'].value
expected['dates'] = [ws_forecast.cell(4 + i, 3).value.strftime('%Y-%m-%d') for i in range(1, exit_p + 1)]
expected['operating_fractions'] = [ws_forecast.cell(4 + i, 5).value for i in range(1, exit_p + 1)]
expected['interest_year_fractions'] = [ws_forecast.cell(4 + i, 6).value for i in range(1, exit_p + 1)]
expected['exit_date'] = ws_returns['B13'].value.strftime('%Y-%m-%d')
expected['holding_period_years'] = ws_returns['B14'].value

# Initial funding
expected['entry_ebitda'] = ws_returns['B5'].value
expected['entry_ev'] = ws_returns['B6'].value
expected['entry_equity'] = ws_returns['B7'].value
expected['ordinary_equity'] = ws_returns['B8'].value
expected['preferred_equity'] = ws_returns['B9'].value
expected['sponsor_investment'] = ws_returns['B10'].value
expected['management_investment'] = ws_returns['B11'].value

# Forecast arrays
expected['revenue'] = [ws_forecast.cell(4 + i, 12).value for i in range(1, exit_p + 1)]
expected['gross_profit'] = [ws_forecast.cell(4 + i, 13).value for i in range(1, exit_p + 1)]
expected['opex'] = [ws_forecast.cell(4 + i, 14).value for i in range(1, exit_p + 1)]
expected['ebitda'] = [ws_forecast.cell(4 + i, 15).value for i in range(1, exit_p + 1)]
expected['ebitda_margin'] = [ws_forecast.cell(4 + i, 16).value for i in range(1, exit_p + 1)]
expected['working_capital'] = [ws_forecast.cell(4 + i, 17).value for i in range(1, exit_p + 1)]
expected['change_working_capital'] = [ws_forecast.cell(4 + i, 18).value for i in range(1, exit_p + 1)]
expected['cash_capex'] = [ws_forecast.cell(4 + i, 19).value for i in range(1, exit_p + 1)]
expected['cash_ebitda'] = [ws_forecast.cell(4 + i, 20).value for i in range(1, exit_p + 1)]

# Financing arrays
expected['cash_interest'] = [ws_debt.cell(4 + i, 8).value for i in range(1, exit_p + 1)]
expected['tlb_interest'] = [ws_debt.cell(4 + i, 6).value for i in range(1, exit_p + 1)]
expected['rcf_interest'] = [ws_debt.cell(4 + i, 7).value for i in range(1, exit_p + 1)]
expected['cash_taxes'] = [ws_debt.cell(4 + i, 10).value for i in range(1, exit_p + 1)]
expected['pre_financing_cash_flow'] = [ws_debt.cell(4 + i, 11).value for i in range(1, exit_p + 1)]
expected['tlb_pik'] = [ws_debt.cell(4 + i, 9).value for i in range(1, exit_p + 1)]
expected['mandatory_tlb_repayment'] = [ws_debt.cell(4 + i, 12).value for i in range(1, exit_p + 1)]
expected['mandatory_rcf_repayment'] = [ws_debt.cell(4 + i, 13).value for i in range(1, exit_p + 1)]
expected['rcf_draw'] = [ws_debt.cell(4 + i, 19).value for i in range(1, exit_p + 1)]
expected['rcf_repayment'] = [ws_debt.cell(4 + i, 23).value for i in range(1, exit_p + 1)]
expected['tlb_sweep'] = [ws_debt.cell(4 + i, 25).value for i in range(1, exit_p + 1)]
expected['tlb_balance'] = [ws_debt.cell(4 + i, 26).value for i in range(1, exit_p + 1)]
expected['rcf_balance'] = [ws_debt.cell(4 + i, 27).value for i in range(1, exit_p + 1)]
expected['cash_balance'] = [ws_debt.cell(4 + i, 28).value for i in range(1, exit_p + 1)]

# Exit bridge
expected['exit_ev'] = ws_returns['B20'].value
expected['exit_fees'] = ws_returns['B21'].value
expected['exit_debt'] = ws_returns['B22'].value
expected['exit_cash'] = ws_returns['B23'].value
expected['equity_value_before_floor'] = ws_returns['B24'].value
expected['equity_proceeds'] = ws_returns['B25'].value

# Waterfall and returns
expected['preferred_accrued'] = ws_returns['B27'].value
expected['preferred_payout'] = ws_returns['B28'].value
expected['ordinary_payout'] = ws_returns['B29'].value
expected['sponsor_proceeds'] = ws_returns['B30'].value
expected['management_proceeds'] = ws_returns['B31'].value
expected['deal_mom'] = ws_returns['B33'].value
expected['sponsor_mom'] = ws_returns['B34'].value
expected['management_mom'] = ws_returns['B35'].value
expected['deal_irr'] = ws_returns['B37'].value
expected['sponsor_irr'] = ws_returns['B38'].value
expected['management_irr'] = ws_returns['B39'].value

print("Total expected fields:", len(expected))
for k, v in list(expected.items())[:15]:
    print(f"  {k}: {v}")
