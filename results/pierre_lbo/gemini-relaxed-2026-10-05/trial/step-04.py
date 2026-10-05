import openpyxl

wb_f = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=False)
wb_v = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=True)

print("=== DEBT COLUMNS ===")
ws_f = wb_f['Debt']
ws_v = wb_v['Debt']
headers_f = [ws_f.cell(4, c).value for c in range(1, ws_f.max_column + 1)]
for idx, h in enumerate(headers_f, 1):
    print(f"Col {idx}: {h}")

print("\n--- Row 5 (Period 1) ---")
for c in range(1, ws_f.max_column + 1):
    f = ws_f.cell(5, c).value
    v = ws_v.cell(5, c).value
    print(f"Col {c:2d} ({headers_f[c-1]}): {f} ==> {v}")

print("\n--- Row 6 (Period 2) ---")
for c in range(1, ws_f.max_column + 1):
    f = ws_f.cell(6, c).value
    v = ws_v.cell(6, c).value
    print(f"Col {c:2d} ({headers_f[c-1]}): {f} ==> {v}")
