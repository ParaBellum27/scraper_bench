import json
import datetime
import math
import openpyxl

wb = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=True)
ws_forecast = wb['Forecast']
ws_debt = wb['Debt']
ws_returns = wb['Returns']
ws_inputs = wb['Inputs']

exit_p = ws_inputs['B8'].value

with open('inputs/base_case.json') as f:
    inputs = json.load(f)

# Run model from previous snippet
# Let's import or paste run_model:
exec('''def edate_12(dt: datetime.date) -> datetime.date:
    y = dt.year + 1
    m = dt.month
    d = dt.day
    if m == 2 and d == 29:
        is_leap = (y % 4 == 0 and (y % 100 != 0 or y % 400 == 0))
        if not is_leap:
            d = 28
    return datetime.date(y, m, d)

def run_model(inputs):
    case_idx = inputs['case'] - 1
    profile = inputs['scenarios'][case_idx]
    
    last_fye = datetime.date.fromisoformat(inputs['last_fiscal_year_end'])
    deal_dt = datetime.date.fromisoformat(inputs['deal_date'])
    exit_period = inputs['exit_period']
    
    dates = []
    prev_fye = last_fye
    operating_fractions = []
    interest_year_fractions = []
    
    for t in range(1, 9):
        curr_fye = edate_12(prev_fye)
        cf_start = max(prev_fye, deal_dt)
        op_frac = (curr_fye - cf_start).days / (curr_fye - prev_fye).days
        int_frac = (curr_fye - cf_start).days / 365.0
        
        dates.append(curr_fye)
        operating_fractions.append(op_frac)
        interest_year_fractions.append(int_frac)
        prev_fye = curr_fye
        
    exit_date_obj = dates[exit_period - 1]
    exit_date = exit_date_obj.isoformat()
    holding_period_years = (exit_date_obj - deal_dt).days / 365.0
    
    rev_growth = profile['revenue_growth']
    gp_margin = profile['gross_profit_margin']
    op_margin = profile['opex_margin']
    capex_in = profile['capex']
    wc_ratio = inputs['working_capital_ratio']
    
    revenue = []
    gross_profit = []
    opex = []
    ebitda = []
    ebitda_margin = []
    working_capital = []
    change_working_capital = []
    cash_capex = []
    cash_ebitda = []
    
    for t in range(8):
        if t == 0:
            rev = inputs['opening_revenue'] * (1.0 + rev_growth[0])
        else:
            rev = revenue[t - 1] * (1.0 + rev_growth[t])
        revenue.append(rev)
        
        gp = rev * gp_margin[t]
        gross_profit.append(gp)
        
        op = rev * op_margin[t]
        opex.append(op)
        
        eb = gp - op
        ebitda.append(eb)
        
        eb_m = gp_margin[t] - op_margin[t]
        ebitda_margin.append(eb_m)
        
        wc = rev * wc_ratio[t]
        working_capital.append(wc)
        
        if t == 0:
            cwc = wc - inputs['opening_working_capital']
        else:
            cwc = wc - working_capital[t - 1]
        change_working_capital.append(cwc)
        
        cash_capex.append(capex_in[t] * operating_fractions[t])
        cash_ebitda.append(eb * operating_fractions[t])

    entry_ebitda = ebitda[0]
    entry_ev = entry_ebitda * inputs['entry_multiple']
    entry_equity = entry_ev - inputs['term_loan']['principal']
    ordinary_equity = entry_equity * inputs['ordinary_equity_fraction']
    preferred_equity = entry_equity - ordinary_equity
    management_investment = ordinary_equity * inputs['management_ordinary_fraction']
    sponsor_investment = entry_equity - management_investment
    
    tl = inputs['term_loan']
    rcf = inputs['revolver']
    
    cash_interest = []
    tlb_interest = []
    rcf_interest = []
    cash_taxes = []
    pre_financing_cash_flow = []
    tlb_pik = []
    mandatory_tlb_repayment = []
    mandatory_rcf_repayment = []
    rcf_draw = []
    rcf_repayment = []
    tlb_sweep = []
    tlb_balance = []
    rcf_balance = []
    cash_balance = []
    
    status = "ok"
    failure_period = None
    funding_shortfall = 0.0
    
    for t in range(1, exit_period + 1):
        idx = t - 1
        if t == 1:
            opening_tlb = float(tl['principal'])
            opening_rcf = 0.0
            opening_cash = 0.0
        else:
            opening_tlb = tlb_balance[-1]
            opening_rcf = rcf_balance[-1]
            opening_cash = cash_balance[-1]
            
        tlb_int = opening_tlb * tl['cash_rate'] * interest_year_fractions[idx]
        rcf_int = opening_rcf * rcf['cash_rate'] * interest_year_fractions[idx]
        c_int = tlb_int + rcf_int
        
        t_pik = opening_tlb * tl['pik_rate'] * interest_year_fractions[idx]
        c_taxes = max(cash_ebitda[idx] - c_int, 0.0) * inputs['tax_rate']
        pf_cf = cash_ebitda[idx] - cash_capex[idx] - change_working_capital[idx] - c_taxes
        
        if t >= tl['maturity_period']:
            prop_mand_tlb = opening_tlb + t_pik
        else:
            prop_mand_tlb = min(opening_tlb + t_pik, tl['principal'] * tl['amortization_rate'] * operating_fractions[idx])
            
        if t >= rcf['maturity_period']:
            prop_mand_rcf = opening_rcf
        else:
            prop_mand_rcf = 0.0
            
        cash_before_rep = opening_cash + pf_cf - c_int
        req_draw = max(prop_mand_tlb + prop_mand_rcf - cash_before_rep, 0.0)
        
        if t < rcf['maturity_period']:
            avail_draw = max(rcf['commitment'] - opening_rcf, 0.0)
        else:
            avail_draw = 0.0
            
        pot_shortfall = max(req_draw - avail_draw, 0.0)
        
        scale = max(opening_cash, abs(cash_ebitda[idx]), cash_capex[idx], abs(change_working_capital[idx]),
                    c_taxes, c_int, prop_mand_tlb + prop_mand_rcf, avail_draw)
        
        if pot_shortfall > 1e-12 * scale:
            status = "liquidity_shortfall"
            failure_period = t
            funding_shortfall = pot_shortfall
            break
            
        draw = min(req_draw, avail_draw)
        cash_after_mand = max(cash_before_rep + draw - prop_mand_tlb - prop_mand_rcf, 0.0)
        interim_rcf = opening_rcf + draw - prop_mand_rcf
        vol_rcf_rep = min(cash_after_mand, interim_rcf)
        total_rcf_rep = prop_mand_rcf + vol_rcf_rep
        cash_after_rcf = cash_after_mand - vol_rcf_rep
        
        sweep = min(cash_after_rcf * tl['cash_sweep'], opening_tlb + t_pik - prop_mand_tlb)
        new_tlb = opening_tlb + t_pik - prop_mand_tlb - sweep
        new_rcf = interim_rcf - vol_rcf_rep
        new_cash = cash_after_rcf - sweep
        
        tlb_interest.append(tlb_int)
        rcf_interest.append(rcf_int)
        cash_interest.append(c_int)
        tlb_pik.append(t_pik)
        cash_taxes.append(c_taxes)
        pre_financing_cash_flow.append(pf_cf)
        mandatory_tlb_repayment.append(prop_mand_tlb)
        mandatory_rcf_repayment.append(prop_mand_rcf)
        rcf_draw.append(draw)
        rcf_repayment.append(total_rcf_rep)
        tlb_sweep.append(sweep)
        tlb_balance.append(new_tlb)
        rcf_balance.append(new_rcf)
        cash_balance.append(new_cash)
        
    if status == "ok":
        N = exit_period - 1
        exit_ev = ebitda[N] * inputs['exit_multiple']
        exit_fees = max(exit_ev, 0.0) * inputs['exit_fee_rate']
        exit_debt = tlb_balance[N] + rcf_balance[N]
        exit_cash = cash_balance[N]
        equity_value_before_floor = exit_ev - exit_fees - exit_debt + exit_cash
        equity_proceeds = max(equity_value_before_floor, 0.0)
        
        pref_accrued = preferred_equity * ((1.0 + inputs['preferred_pik_rate']) ** holding_period_years)
        if equity_proceeds - pref_accrued <= 1e-12 * equity_proceeds:
            pref_payout = equity_proceeds
        else:
            pref_payout = pref_accrued
        ord_payout = equity_proceeds - pref_payout
        
        sponsor_proc = pref_payout + ord_payout * (1.0 - inputs['management_ordinary_fraction'])
        mgmt_proc = ord_payout * inputs['management_ordinary_fraction']
        
        deal_mom = (equity_proceeds / entry_equity) if entry_equity > 0 else None
        sponsor_mom = (sponsor_proc / sponsor_investment) if sponsor_investment > 0 else None
        mgmt_mom = (mgmt_proc / management_investment) if management_investment > 0 else None
        
        deal_irr = ((equity_proceeds / entry_equity) ** (1.0 / holding_period_years) - 1.0) if (entry_equity > 0 and equity_proceeds > 0) else None
        sponsor_irr = ((sponsor_proc / sponsor_investment) ** (1.0 / holding_period_years) - 1.0) if (sponsor_investment > 0 and sponsor_proc > 0) else None
        mgmt_irr = ((mgmt_proc / management_investment) ** (1.0 / holding_period_years) - 1.0) if (management_investment > 0 and mgmt_proc > 0) else None
    else:
        exit_ev = None
        exit_fees = None
        exit_debt = None
        exit_cash = None
        equity_value_before_floor = None
        equity_proceeds = None
        pref_accrued = None
        pref_payout = None
        ord_payout = None
        sponsor_proc = None
        mgmt_proc = None
        deal_mom = None
        sponsor_mom = None
        mgmt_mom = None
        deal_irr = None
        sponsor_irr = None
        mgmt_irr = None
        
    return {
        "status": status, "failure_period": failure_period, "funding_shortfall": funding_shortfall,
        "dates": [d.isoformat() for d in dates[:exit_period]],
        "operating_fractions": operating_fractions[:exit_period],
        "interest_year_fractions": interest_year_fractions[:exit_period],
        "exit_date": exit_date, "holding_period_years": holding_period_years,
        "entry_ebitda": entry_ebitda, "entry_ev": entry_ev, "entry_equity": entry_equity,
        "ordinary_equity": ordinary_equity, "preferred_equity": preferred_equity,
        "sponsor_investment": sponsor_investment, "management_investment": management_investment,
        "revenue": revenue[:exit_period], "gross_profit": gross_profit[:exit_period],
        "opex": opex[:exit_period], "ebitda": ebitda[:exit_period], "ebitda_margin": ebitda_margin[:exit_period],
        "working_capital": working_capital[:exit_period], "change_working_capital": change_working_capital[:exit_period],
        "cash_capex": cash_capex[:exit_period], "cash_ebitda": cash_ebitda[:exit_period],
        "cash_interest": cash_interest, "tlb_interest": tlb_interest, "rcf_interest": rcf_interest,
        "cash_taxes": cash_taxes, "pre_financing_cash_flow": pre_financing_cash_flow, "tlb_pik": tlb_pik,
        "mandatory_tlb_repayment": mandatory_tlb_repayment, "mandatory_rcf_repayment": mandatory_rcf_repayment,
        "rcf_draw": rcf_draw, "rcf_repayment": rcf_repayment, "tlb_sweep": tlb_sweep,
        "tlb_balance": tlb_balance, "rcf_balance": rcf_balance, "cash_balance": cash_balance,
        "exit_ev": exit_ev, "exit_fees": exit_fees, "exit_debt": exit_debt, "exit_cash": exit_cash,
        "equity_value_before_floor": equity_value_before_floor, "equity_proceeds": equity_proceeds,
        "preferred_accrued": pref_accrued, "preferred_payout": pref_payout, "ordinary_payout": ord_payout,
        "sponsor_proceeds": sponsor_proc, "management_proceeds": mgmt_proc,
        "deal_mom": deal_mom, "sponsor_mom": sponsor_mom, "management_mom": mgmt_mom,
        "deal_irr": deal_irr, "sponsor_irr": sponsor_irr, "management_irr": mgmt_irr,
    }
''')

res = run_model(inputs)

# Compare
expected = {
    'status': ws_returns['B16'].value,
    'failure_period': ws_returns['B17'].value if ws_returns['B17'].value != "" else None,
    'funding_shortfall': ws_returns['B18'].value,
    'dates': [ws_forecast.cell(4 + i, 3).value.strftime('%Y-%m-%d') for i in range(1, exit_p + 1)],
    'operating_fractions': [ws_forecast.cell(4 + i, 5).value for i in range(1, exit_p + 1)],
    'interest_year_fractions': [ws_forecast.cell(4 + i, 6).value for i in range(1, exit_p + 1)],
    'exit_date': ws_returns['B13'].value.strftime('%Y-%m-%d'),
    'holding_period_years': ws_returns['B14'].value,
    'entry_ebitda': ws_returns['B5'].value,
    'entry_ev': ws_returns['B6'].value,
    'entry_equity': ws_returns['B7'].value,
    'ordinary_equity': ws_returns['B8'].value,
    'preferred_equity': ws_returns['B9'].value,
    'sponsor_investment': ws_returns['B10'].value,
    'management_investment': ws_returns['B11'].value,
    'revenue': [ws_forecast.cell(4 + i, 12).value for i in range(1, exit_p + 1)],
    'gross_profit': [ws_forecast.cell(4 + i, 13).value for i in range(1, exit_p + 1)],
    'opex': [ws_forecast.cell(4 + i, 14).value for i in range(1, exit_p + 1)],
    'ebitda': [ws_forecast.cell(4 + i, 15).value for i in range(1, exit_p + 1)],
    'ebitda_margin': [ws_forecast.cell(4 + i, 16).value for i in range(1, exit_p + 1)],
    'working_capital': [ws_forecast.cell(4 + i, 17).value for i in range(1, exit_p + 1)],
    'change_working_capital': [ws_forecast.cell(4 + i, 18).value for i in range(1, exit_p + 1)],
    'cash_capex': [ws_forecast.cell(4 + i, 19).value for i in range(1, exit_p + 1)],
    'cash_ebitda': [ws_forecast.cell(4 + i, 20).value for i in range(1, exit_p + 1)],
    'cash_interest': [ws_debt.cell(4 + i, 8).value for i in range(1, exit_p + 1)],
    'tlb_interest': [ws_debt.cell(4 + i, 6).value for i in range(1, exit_p + 1)],
    'rcf_interest': [ws_debt.cell(4 + i, 7).value for i in range(1, exit_p + 1)],
    'cash_taxes': [ws_debt.cell(4 + i, 10).value for i in range(1, exit_p + 1)],
    'pre_financing_cash_flow': [ws_debt.cell(4 + i, 11).value for i in range(1, exit_p + 1)],
    'tlb_pik': [ws_debt.cell(4 + i, 9).value for i in range(1, exit_p + 1)],
    'mandatory_tlb_repayment': [ws_debt.cell(4 + i, 12).value for i in range(1, exit_p + 1)],
    'mandatory_rcf_repayment': [ws_debt.cell(4 + i, 13).value for i in range(1, exit_p + 1)],
    'rcf_draw': [ws_debt.cell(4 + i, 19).value for i in range(1, exit_p + 1)],
    'rcf_repayment': [ws_debt.cell(4 + i, 23).value for i in range(1, exit_p + 1)],
    'tlb_sweep': [ws_debt.cell(4 + i, 25).value for i in range(1, exit_p + 1)],
    'tlb_balance': [ws_debt.cell(4 + i, 26).value for i in range(1, exit_p + 1)],
    'rcf_balance': [ws_debt.cell(4 + i, 27).value for i in range(1, exit_p + 1)],
    'cash_balance': [ws_debt.cell(4 + i, 28).value for i in range(1, exit_p + 1)],
    'exit_ev': ws_returns['B20'].value,
    'exit_fees': ws_returns['B21'].value,
    'exit_debt': ws_returns['B22'].value,
    'exit_cash': ws_returns['B23'].value,
    'equity_value_before_floor': ws_returns['B24'].value,
    'equity_proceeds': ws_returns['B25'].value,
    'preferred_accrued': ws_returns['B27'].value,
    'preferred_payout': ws_returns['B28'].value,
    'ordinary_payout': ws_returns['B29'].value,
    'sponsor_proceeds': ws_returns['B30'].value,
    'management_proceeds': ws_returns['B31'].value,
    'deal_mom': ws_returns['B33'].value,
    'sponsor_mom': ws_returns['B34'].value,
    'management_mom': ws_returns['B35'].value,
    'deal_irr': ws_returns['B37'].value,
    'sponsor_irr': ws_returns['B38'].value,
    'management_irr': ws_returns['B39'].value,
}

all_match = True
for k in expected:
    ev = expected[k]
    rv = res[k]
    if isinstance(ev, list):
        if len(ev) != len(rv):
            print(f"Mismatch in length for {k}: {len(ev)} vs {len(rv)}")
            all_match = False
        else:
            for i, (e_item, r_item) in enumerate(zip(ev, rv)):
                if isinstance(e_item, (int, float)):
                    if abs(e_item - r_item) > 1e-7 * max(abs(e_item), 1.0):
                        print(f"Mismatch in {k}[{i}]: {e_item} vs {r_item}")
                        all_match = False
                elif e_item != r_item:
                    print(f"Mismatch in {k}[{i}]: {e_item} vs {r_item}")
                    all_match = False
    elif isinstance(ev, (int, float)):
        if abs(ev - rv) > 1e-7 * max(abs(ev), 1.0):
            print(f"Mismatch in {k}: {ev} vs {rv}")
            all_match = False
    elif ev != rv:
        print(f"Mismatch in {k}: {ev} vs {rv}")
        all_match = False

if all_match:
    print("ALL 55 FIELDS MATCHED PERFECTLY ON BASE CASE!")
