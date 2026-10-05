import openpyxl

wb_f = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=False)

def dump_sheet_formulas(sheetname):
    ws = wb_f[sheetname]
    print(f"=== {sheetname} ===")
    max_c = ws.max_column
    for r in range(4, 13):
        vals = [f"{openpyxl.utils.get_column_letter(c)}{r}: {ws.cell(r, c).value}" for c in range(1, max_c + 1) if ws.cell(r, c).value is not None]
        print(f"Row {r}: " + "; ".join(vals))

dump_sheet_formulas('Forecast')
dump_sheet_formulas('Debt')
