import openpyxl

wb_f = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=False)
wb_v = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=True)

ws_f = wb_f['Returns']
ws_v = wb_v['Returns']

for r in range(1, ws_f.max_row + 1):
    f_vals = [ws_f.cell(r, c).value for c in range(1, ws_f.max_column + 1)]
    v_vals = [ws_v.cell(r, c).value for c in range(1, ws_v.max_column + 1)]
    if any(x is not None for x in f_vals):
        row_str = " | ".join(f"{str(f)} [{str(v)}]" if f != v else str(f) for f, v in zip(f_vals, v_vals) if f is not None or v is not None)
        print(f"R{r}: {row_str}")
