import openpyxl
import math

wb = openpyxl.load_workbook('pierre_lbo_one.xlsx', data_only=True)
ws = wb['Returns']

for name, r, init_cell, proc_cell in [
    ("deal", 37, 'B7', 'B25'),
    ("sponsor", 38, 'B10', 'B30'),
    ("management", 39, 'B11', 'B31'),
]:
    xirr_val = ws[f'B{r}'].value
    i0 = ws[init_cell].value
    p = ws[proc_cell].value
    h = ws['B14'].value
    cagr = (p / i0) ** (1 / h) - 1
    exp_ln = math.exp(math.log(p / i0) / h) - 1
    print(f"{name}:")
    print(f"  Workbook XIRR: {xirr_val:.15f}")
    print(f"  Closed form  : {cagr:.15f}")
    print(f"  Exp(Ln)      : {exp_ln:.15f}")
    print(f"  Diff         : {abs(xirr_val - cagr):.2e}")
