import openpyxl

wb = openpyxl.load_workbook('fcff2st.xlsx', data_only=True)
sheet = wb['NewFCFF2Stage']

# Check the flags for F99, F100
print("=== Flags ===")
print(f"F99: {sheet.cell(99, 6).value}")  # F = column 6
print(f"F100: {sheet.cell(100, 6).value}")
print(f"F101: {sheet.cell(101, 6).value}")
print(f"F102: {sheet.cell(102, 6).value}")

# Also check what these questions are
wb_formulas = openpyxl.load_workbook('fcff2st.xlsx', data_only=False)
sheet_formulas = wb_formulas['NewFCFF2Stage']

for row in range(99, 103):
    for col in range(1, 8):
        cell = sheet_formulas.cell(row, col)
        if cell.value:
            print(f"R{row}C{col}: {cell.value}")

# Now check N135 value
print(f"\nN135 value: {sheet.cell(135, 14).value}")

# And N136
print(f"N136 value: {sheet.cell(136, 14).value}")

# Let me also check what D23 is
print(f"D23 (current depreciation): {sheet.cell(23, 4).value}")
print(f"D128 (growth rate depreciation): {sheet.cell(128, 4).value}")
print(f"D129 (growth rate revenues): {sheet.cell(129, 4).value}")
print(f"E131 (working capital %): {sheet.cell(131, 5).value}")

# Check F99, F100 values as strings
print(f"\nF99 as string: {str(sheet_formulas.cell(99, 6).value)}")
print(f"F100 as string: {str(sheet_formulas.cell(100, 6).value)}")
print(f"F101 as string: {str(sheet_formulas.cell(101, 6).value)}")
print(f"F102 as string: {str(sheet_formulas.cell(102, 6).value)}")