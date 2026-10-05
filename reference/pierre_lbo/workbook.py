"""Independent, formula-only Pierre LBO One Excel derivative.

Build (does not calculate or launch Excel)::

    python reference/pierre_lbo/workbook.py build \
        --inputs tasks/pierre_lbo/inputs/base_case.json \
        --output runs/private/pierre_lbo/pierre_lbo_one.xlsx

Recalculate the derivative's sheets in Microsoft Excel in their tab order, save,
and close it before extracting its cached results::

    python reference/pierre_lbo/workbook.py extract \
        --workbook runs/private/pierre_lbo/pierre_lbo_one.xlsx

Extraction writes the exact 55-field JSON object to stdout. Formula caches are
required: openpyxl alone cannot calculate this model. No Python valuation is
called, and no calculated expected outputs are stored in the generated file.

Workbook map (all sheets visible):
* Inputs!A5:C26: scalar input paths, editable values and descriptions;
  Inputs!A31:B38: eight working-capital ratios.
* Scenarios!A5:F36: four eight-period scenario profiles, in case/period order.
* Forecast!A5:T12: annual operating drivers, fiscal dates and stub cash flows.
* Debt!A5:AB12: sequential funding tests and executed financing. Candidate
  obligations may be visible on the failed row, but executed flows/balances
  are blank. No subsequent financing row is evaluated.
* Returns!B5:B11: entry funding; B13:B14: planned exit timing;
  B16:B18: funding status; B20:B25: exit; B27:B31: priority waterfall;
  B33:B35: investment multiples; B37:B39: native Excel XIRR results.
  Returns!B44:E45: dated investment and exit flows used by XIRR.

Null outputs are represented by formula-generated empty strings. Source and
frozen workbooks are never legitimate build destinations. The generated
workbook has no macros, links, data tables, hidden sheets or reference code.
"""

from __future__ import annotations

import argparse
import calendar
from datetime import date, datetime
import json
import math
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties


# Stable, public cell mappings; extraction reads Excel caches at these addresses.
_INPUTS = (
    ("case", "Scenario selector: integer 1–4"),
    ("last_fiscal_year_end", "ISO date; successive fiscal ends use EDATE(previous,12)"),
    ("deal_date", "ISO closing date, strictly within the first fiscal period"),
    ("exit_period", "Forecast period 1–8; first post-close period is 1"),
    ("opening_revenue", "Annual entry assumption, GBP thousands"),
    ("opening_working_capital", "Entry working-capital balance, GBP thousands"),
    ("entry_multiple", "Multiple of first annual forecast EBITDA"),
    ("exit_multiple", "Multiple of exit-period annual EBITDA"),
    ("exit_fee_rate", "Fraction of nonnegative exit enterprise value"),
    ("tax_rate", "Fraction of positive cash EBITDA less cash interest"),
    ("ordinary_equity_fraction", "Fraction of entry equity invested in ordinary shares"),
    ("management_ordinary_fraction", "Management ownership of ordinary shares only"),
    ("preferred_pik_rate", "Effective annual preferred accrual, ACT/365"),
    ("term_loan.principal", "Entry term loan, GBP thousands"),
    ("term_loan.cash_rate", "Annual opening-balance cash interest, ACT/365"),
    ("term_loan.pik_rate", "Annual opening-balance PIK, ACT/365"),
    ("term_loan.amortization_rate", "Fraction of original principal, operating-period prorated"),
    ("term_loan.cash_sweep", "Fraction of cash remaining after all revolver repayments"),
    ("term_loan.maturity_period", "Fiscal period bucket; full TLB plus PIK becomes due"),
    ("revolver.commitment", "Maximum outstanding revolver, GBP thousands"),
    ("revolver.cash_rate", "Annual opening-balance cash interest, ACT/365"),
    ("revolver.maturity_period", "Fiscal period bucket; no draws at or after maturity"),
)
_INPUT_ROWS = {key: row for row, (key, _) in enumerate(_INPUTS, 5)}
_FUNDING_ROWS = {
    "entry_ebitda": 5, "entry_ev": 6, "entry_equity": 7,
    "ordinary_equity": 8, "preferred_equity": 9,
    "sponsor_investment": 10, "management_investment": 11,
}
_FORECAST_COLUMNS = {
    "revenue": "L", "gross_profit": "M", "opex": "N", "ebitda": "O",
    "ebitda_margin": "P", "working_capital": "Q", "change_working_capital": "R",
    "cash_capex": "S", "cash_ebitda": "T",
}
_FINANCING_COLUMNS = {
    "cash_interest": "H", "tlb_interest": "F", "rcf_interest": "G",
    "cash_taxes": "J", "pre_financing_cash_flow": "K", "tlb_pik": "I",
    "mandatory_tlb_repayment": "L", "mandatory_rcf_repayment": "M",
    "rcf_draw": "S", "rcf_repayment": "W", "tlb_sweep": "Y",
    "tlb_balance": "Z", "rcf_balance": "AA", "cash_balance": "AB",
}
_EXIT_ROWS = {
    "exit_ev": 20, "exit_fees": 21, "exit_debt": 22, "exit_cash": 23,
    "equity_value_before_floor": 24, "equity_proceeds": 25,
    "preferred_accrued": 27, "preferred_payout": 28, "ordinary_payout": 29,
    "sponsor_proceeds": 30, "management_proceeds": 31,
    "deal_mom": 33, "sponsor_mom": 34, "management_mom": 35,
    "deal_irr": 37, "sponsor_irr": 38, "management_irr": 39,
}
_SCENARIO_FIELDS = ("revenue_growth", "gross_profit_margin", "opex_margin", "capex")
_NUMBER_FORMAT = '#,##0.000000;[Red](#,##0.000000);0.000000'
_PERCENT_FORMAT = '0.000000%'


def _input(key: str) -> str:
    return f"Inputs!$B${_INPUT_ROWS[key]}"


def _lookup(inputs: dict, key: str) -> Any:
    value = inputs
    for part in key.split("."):
        value = value[part]
    return value


def _validate_inputs(inputs: dict) -> None:
    """Validate the input contract, without deriving any workbook output caches."""
    def fields(value: Any, expected: set[str], name: str) -> dict[str, Any]:
        if not isinstance(value, dict) or set(value) != expected:
            raise ValueError(f"{name} must have exactly these keys: {sorted(expected)}")
        return value

    def numeric(value: Any, name: str, low: float | None = 0,
                high: float | None = None, strict: bool = False) -> None:
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise ValueError(f"{name} must be a finite number, not a boolean")
        try:
            finite = math.isfinite(value)
        except OverflowError:
            finite = False
        if not finite:
            raise ValueError(f"{name} must be finite")
        if low is not None and (value <= low if strict else value < low):
            raise ValueError(f"{name} is below its allowed domain")
        if high is not None and value > high:
            raise ValueError(f"{name} exceeds {high}")

    def integer(value: Any, name: str, high: int | None = None) -> None:
        if isinstance(value, bool) or not isinstance(value, int):
            raise ValueError(f"{name} must be an integer")
        numeric(value, name, 1, high)

    expected = {key.split(".")[0] for key, _ in _INPUTS}
    fields(inputs, expected | {"working_capital_ratio", "scenarios"}, "inputs")
    fields(inputs["term_loan"], {"principal", "cash_rate", "pik_rate",
           "amortization_rate", "cash_sweep", "maturity_period"}, "term_loan")
    fields(inputs["revolver"], {"commitment", "cash_rate", "maturity_period"}, "revolver")
    integer(inputs["case"], "case", 4)
    integer(inputs["exit_period"], "exit_period", 8)
    for key in ("term_loan.maturity_period", "revolver.maturity_period"):
        integer(_lookup(inputs, key), key)
    dates = {}
    for key in ("last_fiscal_year_end", "deal_date"):
        value = inputs[key]
        if not isinstance(value, str):
            raise ValueError(f"{key} must be an ISO date string")
        try:
            parsed = date.fromisoformat(value)
        except ValueError as exc:
            raise ValueError(f"{key} must be an ISO date string") from exc
        if parsed.isoformat() != value:
            raise ValueError(f"{key} must use YYYY-MM-DD")
        dates[key] = parsed
    last = dates["last_fiscal_year_end"]
    try:
        first = date(last.year + 1, last.month,
                     min(last.day, calendar.monthrange(last.year + 1, last.month)[1]))
    except ValueError as exc:
        raise ValueError("First fiscal year-end is outside the supported date range") from exc
    if not last < dates["deal_date"] < first:
        raise ValueError("deal_date must be strictly between last and next fiscal year-end")
    unit_fractions = {"exit_fee_rate", "tax_rate", "ordinary_equity_fraction",
                      "management_ordinary_fraction", "term_loan.amortization_rate",
                      "term_loan.cash_sweep"}
    skip = {"case", "exit_period", "deal_date", "last_fiscal_year_end",
            "term_loan.maturity_period", "revolver.maturity_period"}
    for key, _ in _INPUTS:
        if key not in skip:
            numeric(_lookup(inputs, key), key, 0, 1 if key in unit_fractions else None,
                    strict=key == "opening_revenue")

    def vector(value: Any, name: str, low: float, high: float | None = None,
               strict: bool = False) -> None:
        if not isinstance(value, list) or len(value) != 8:
            raise ValueError(f"{name} must be an eight-element array")
        for index, item in enumerate(value):
            numeric(item, f"{name}[{index}]", low, high, strict)

    vector(inputs["working_capital_ratio"], "working_capital_ratio", 0)
    if not isinstance(inputs["scenarios"], list) or len(inputs["scenarios"]) != 4:
        raise ValueError("scenarios must contain exactly four profiles")
    for index, scenario in enumerate(inputs["scenarios"]):
        scenario = fields(scenario, set(_SCENARIO_FIELDS), f"scenarios[{index}]")
        for key in _SCENARIO_FIELDS:
            vector(scenario[key], f"scenarios[{index}].{key}",
                   -1 if key == "revenue_growth" else 0,
                   1 if key in ("gross_profit_margin", "opex_margin") else None,
                   strict=key == "revenue_growth")
    scenario = inputs["scenarios"][inputs["case"] - 1]
    annual_revenue = inputs["opening_revenue"] * (1 + scenario["revenue_growth"][0])
    annual_ebitda = (annual_revenue * scenario["gross_profit_margin"][0]
                     - annual_revenue * scenario["opex_margin"][0])
    entry_equity = annual_ebitda * inputs["entry_multiple"] - inputs["term_loan"]["principal"]
    if not math.isfinite(entry_equity) or entry_equity <= 0:
        raise ValueError("Derived entry_equity must be finite and strictly positive")


def _sheet(workbook: Workbook, name: str, title: str, headers: list[str]):
    sheet = workbook.create_sheet(name)
    sheet["A1"] = title
    sheet["A1"].font = Font(size=16, bold=True, color="17365D")
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(4, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(column)].width = 22
    sheet.row_dimensions[4].height = 44
    sheet.freeze_panes = "B5"
    sheet.sheet_view.showGridLines = False
    return sheet


def _formula(sheet, address: str, expression: str) -> None:
    cell = sheet[address]
    cell.value = "=" + expression
    cell.number_format = _NUMBER_FORMAT
    cell.font = Font(color="000000")


def build_workbook(inputs: dict, destination: Path) -> None:
    """Write an uncalculated, independently formulated .xlsx derivative."""
    _validate_inputs(inputs)
    destination = Path(destination).expanduser().resolve()
    protected = Path(__file__).resolve().parents[2] / "runs/private/pierre_lbo/source"
    if (destination == protected or protected in destination.parents
            or destination.name.lower() == "pierre_lbo.xlsx"
            or any(part.lower() == "frozen" for part in destination.parts)):
        raise ValueError("Refusing to write a source or frozen workbook destination")
    if destination.exists():
        raise FileExistsError("Refusing to overwrite an existing workbook; choose a new derivative path")
    if destination.suffix.lower() != ".xlsx":
        raise ValueError("The derivative destination must have an .xlsx extension")

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.calculation = CalcProperties(calcId=0, calcMode="auto", fullCalcOnLoad=True,
                                          forceFullCalc=True, calcOnSave=True)
    workbook.properties.title = "Pierre LBO One"
    workbook.properties.subject = "Independent corrected leveraged buyout derivative"
    workbook.properties.creator = "Pierre LBO One"
    workbook.properties.description = "Explicit Excel formulas; GBP thousands; ACT/365 interest."
    sheet = _sheet(workbook, "Inputs", "Pierre LBO One — inputs", ["Input path", "Value", "Definition"])
    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["C"].width = 88
    for key, description in _INPUTS:
        row = _INPUT_ROWS[key]
        sheet.cell(row, 1, key)
        cell = sheet.cell(row, 2, _lookup(inputs, key))
        cell.font = Font(color="0000FF")
        cell.number_format = _NUMBER_FORMAT
        if key in ("last_fiscal_year_end", "deal_date"):
            cell.value = date.fromisoformat(cell.value)
            cell.number_format = "yyyy-mm-dd"
        elif key in ("case", "exit_period") or key.endswith("maturity_period"):
            cell.number_format = "0"
        elif key.endswith(("rate", "fraction", "cash_sweep")):
            cell.number_format = _PERCENT_FORMAT
        sheet.cell(row, 3, description)
    sheet["A30"], sheet["B30"] = "Forecast period", "working_capital_ratio"
    for period, ratio in enumerate(inputs["working_capital_ratio"], 1):
        sheet.cell(period + 30, 1, period)
        cell = sheet.cell(period + 30, 2, ratio)
        cell.number_format = _PERCENT_FORMAT
        cell.font = Font(color="0000FF")
    notes = [
        "All monetary amounts are GBP thousands. Blue cells are inputs; black cells are formulas.",
        "Input schema: scalar paths above; working_capital_ratio[8]; scenarios[4] of four 8-element arrays.",
        "Scenarios columns: revenue_growth, gross_profit_margin, opex_margin, capex (positive expenditure).",
        "Opening cash and revolver are zero. Entry is funded only by the term loan and equity; no entry fees.",
        "Annual EBITDA values entry/exit. Cash EBITDA and capex are stubbed; working-capital changes are not.",
        "No D&A, tax losses/refunds or deduction of PIK/preferred accrual; tax is on positive cash EBITDA less cash interest.",
        "Interest uses opening debt and ACT/365. Fiscal dates preserve the input day with sequential EDATE(previous,12).",
        "Revolver is repaid before term-loan cash sweep. Maturity is a fiscal period bucket, not a legal anniversary.",
        "Financing stops before an infeasible row. Forecast remains planned; exit/return cells become blank (JSON null).",
        "Preferred belongs only to the sponsor. No interim dividends; zero investment or zero payout has no finite XIRR.",
        "At the preferred hurdle, residual <= 1E-12 of equity proceeds is roundoff: all proceeds go to preferred.",
        "XIRR uses actual dated cash flows; a log-ratio starting guess stabilizes convergence for extreme two-flow returns.",
        "Recalculate and save in Microsoft Excel before extracting cached outputs. No external links, macros or data tables.",
    ]
    for row, note in enumerate(notes, 41):
        sheet.cell(row, 1, note)

    sheet = _sheet(workbook, "Scenarios", "Scenario profiles — annual operating assumptions",
                   ["Case", "Period", *_SCENARIO_FIELDS])
    for scenario_index, scenario in enumerate(inputs["scenarios"]):
        for period in range(8):
            row = 5 + scenario_index * 8 + period
            sheet.cell(row, 1, scenario_index + 1)
            sheet.cell(row, 2, period + 1)
            for column, key in enumerate(_SCENARIO_FIELDS, 3):
                cell = sheet.cell(row, column, scenario[key][period])
                cell.font = Font(color="0000FF")
                cell.number_format = _NUMBER_FORMAT if key == "capex" else _PERCENT_FORMAT
    sheet.auto_filter.ref = "A4:F36"

    sheet = _sheet(workbook, "Forecast", "Annual forecast and actual-day cash-flow timing", [
        "Period", "Previous fiscal end", "Fiscal end / dates", "Cash-flow start",
        "operating_fractions", "interest_year_fractions", "Selected revenue growth",
        "Selected GP margin", "Selected opex margin", "Annual capex", "WC ratio",
        "revenue", "gross_profit", "opex", "ebitda", "ebitda_margin",
        "working_capital", "change_working_capital", "cash_capex", "cash_ebitda",
    ])
    for period in range(1, 9):
        row = period + 4
        sheet.cell(row, 1, period)
        prior_revenue = _input("opening_revenue") if period == 1 else f"L{row - 1}"
        prior_wc = _input("opening_working_capital") if period == 1 else f"Q{row - 1}"
        expressions = {
            "B": _input("last_fiscal_year_end") if period == 1 else f"C{row - 1}",
            "C": f"EDATE(B{row},12)", "D": f"MAX(B{row},{_input('deal_date')})",
            "E": f"(C{row}-D{row})/(C{row}-B{row})", "F": f"(C{row}-D{row})/365",
            "K": f"Inputs!B{period + 30}", "L": f"{prior_revenue}*(1+G{row})",
            "M": f"L{row}*H{row}", "N": f"L{row}*I{row}", "O": f"M{row}-N{row}",
            "P": f"H{row}-I{row}", "Q": f"L{row}*K{row}", "R": f"Q{row}-{prior_wc}",
            "S": f"J{row}*E{row}", "T": f"O{row}*E{row}",
        }
        for target, source in zip(("G", "H", "I", "J"), ("C", "D", "E", "F")):
            expressions[target] = f"INDEX(Scenarios!${source}$5:${source}$36,({_input('case')}-1)*8+A{row})"
        for column, expression in expressions.items():
            _formula(sheet, f"{column}{row}", expression)
        for column in ("B", "C", "D"):
            sheet[f"{column}{row}"].number_format = "yyyy-mm-dd"
        for column in ("E", "F", "G", "H", "I", "K", "P"):
            sheet[f"{column}{row}"].number_format = _PERCENT_FORMAT

    sheet = _sheet(workbook, "Debt", "Sequential debt, cash and liquidity waterfall", [
        "Period", "Evaluate row?", "Opening TLB", "Opening RCF", "Opening cash",
        "tlb_interest", "rcf_interest", "cash_interest", "tlb_pik", "cash_taxes",
        "pre_financing_cash_flow", "Proposed mandatory TLB repayment",
        "Proposed mandatory RCF repayment", "Cash before repayments", "Required draw",
        "Available draw", "Potential funding shortfall", "Row status", "rcf_draw",
        "Cash after mandatory repayments", "Interim RCF", "Voluntary RCF repayment",
        "rcf_repayment (total)", "Cash after RCF repayment", "tlb_sweep", "tlb_balance",
        "rcf_balance", "cash_balance",
    ])
    sheet["A2"] = "Candidate obligations are diagnostic only on a failed row; executed flows and balances are blank."
    for period in range(1, 9):
        row = period + 4
        sheet.cell(row, 1, period)
        previous_ok = "TRUE" if period == 1 else f'R{row - 1}="ok"'
        _formula(sheet, f"B{row}", f"AND(A{row}<={_input('exit_period')},{previous_ok})")
        candidates = {
            "C": _input("term_loan.principal") if period == 1 else f"Z{row - 1}",
            "D": "0" if period == 1 else f"AA{row - 1}",
            "E": "0" if period == 1 else f"AB{row - 1}",
            "F": f"C{row}*{_input('term_loan.cash_rate')}*Forecast!F{row}",
            "G": f"D{row}*{_input('revolver.cash_rate')}*Forecast!F{row}",
            "H": f"F{row}+G{row}",
            "I": f"C{row}*{_input('term_loan.pik_rate')}*Forecast!F{row}",
            "J": f"MAX(Forecast!T{row}-H{row},0)*{_input('tax_rate')}",
            "K": f"Forecast!T{row}-Forecast!S{row}-Forecast!R{row}-J{row}",
            "L": (f"IF(A{row}>={_input('term_loan.maturity_period')},C{row}+I{row},"
                  f"MIN(C{row}+I{row},{_input('term_loan.principal')}*"
                  f"{_input('term_loan.amortization_rate')}*Forecast!E{row}))"),
            "M": f"IF(A{row}>={_input('revolver.maturity_period')},D{row},0)",
            "N": f"E{row}+K{row}-H{row}", "O": f"MAX(L{row}+M{row}-N{row},0)",
            "P": (f"IF(A{row}<{_input('revolver.maturity_period')},"
                  f"MAX({_input('revolver.commitment')}-D{row},0),0)"),
            "Q": f"MAX(O{row}-P{row},0)",
        }
        for column, expression in candidates.items():
            _formula(sheet, f"{column}{row}", f'IF(B{row},{expression},"")')
        _formula(sheet, f"R{row}",
                 f'IF(B{row},IF(Q{row}>1E-8,"liquidity_shortfall","ok"),"not_projected")')
        executed = {
            "S": f"MIN(O{row},P{row})", "T": f"MAX(N{row}+S{row}-L{row}-M{row},0)",
            "U": f"D{row}+S{row}-M{row}", "V": f"MIN(T{row},U{row})",
            "W": f"M{row}+V{row}", "X": f"T{row}-V{row}",
            "Y": f"MIN(X{row}*{_input('term_loan.cash_sweep')},C{row}+I{row}-L{row})",
            "Z": f"C{row}+I{row}-L{row}-Y{row}", "AA": f"U{row}-V{row}",
            "AB": f"X{row}-Y{row}",
        }
        for column, expression in executed.items():
            _formula(sheet, f"{column}{row}", f'IF(R{row}="ok",{expression},"")')

    sheet = _sheet(workbook, "Returns", "Entry funding, exit priority waterfall and native XIRR", ["Output", "Value"])
    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["B"].width = 28
    expressions = {
        5: "Forecast!O5", 6: f"B5*{_input('entry_multiple')}",
        7: f"B6-{_input('term_loan.principal')}",
        8: f"B7*{_input('ordinary_equity_fraction')}", 9: "B7-B8", 10: "B7-B11",
        11: f"B8*{_input('management_ordinary_fraction')}",
        13: f"INDEX(Forecast!C5:C12,{_input('exit_period')})",
        14: f"(B13-{_input('deal_date')})/365",
        16: 'IF(COUNTIF(Debt!R5:R12,"liquidity_shortfall")>0,"liquidity_shortfall","ok")',
        17: 'IF(B16="ok","",MATCH("liquidity_shortfall",Debt!R5:R12,0))',
        18: 'IF(B16="ok",0,INDEX(Debt!Q5:Q12,B17))',
    }
    exit_expressions = {
        20: f"INDEX(Forecast!O5:O12,{_input('exit_period')})*{_input('exit_multiple')}",
        21: f"MAX(B20,0)*{_input('exit_fee_rate')}",
        22: (f"INDEX(Debt!Z5:Z12,{_input('exit_period')})"
             f"+INDEX(Debt!AA5:AA12,{_input('exit_period')})"),
        23: f"INDEX(Debt!AB5:AB12,{_input('exit_period')})", 24: "B20-B21-B22+B23",
        25: "MAX(B24,0)", 27: f"B9*(1+{_input('preferred_pik_rate')})^B14",
        28: "IF(B25-B27<=1E-12*B25,B25,B27)", 29: "B25-B28",
        30: f"B28+B29*(1-{_input('management_ordinary_fraction')})",
        31: f"B29*{_input('management_ordinary_fraction')}",
    }
    expressions.update({row: f'IF($B$16="ok",{expression},"")'
                        for row, expression in exit_expressions.items()})
    for mom_row, irr_row, investment_row, payout_row, flow_column in (
        (33, 37, 7, 25, "C"), (34, 38, 10, 30, "D"), (35, 39, 11, 31, "E"),
    ):
        expressions[mom_row] = (f'IF(AND($B$16="ok",B{investment_row}>0),'
                                f'B{payout_row}/B{investment_row},"")')
        # The native solver remains the source of the return. A near-root guess
        # prevents default-guess convergence failures for highly impaired exits.
        guess = f"MAX(-0.999999999999999,EXP(LN(B{payout_row}/B{investment_row})/$B$14)-1)"
        expressions[irr_row] = (
            f'IF(AND($B$16="ok",B{investment_row}>0,B{payout_row}>0),'
            f'XIRR({flow_column}44:{flow_column}45,$B$44:$B$45,{guess}),"")'
        )
        _formula(sheet, f"{flow_column}44", f"-B{investment_row}")
        _formula(sheet, f"{flow_column}45", f'IF($B$16="ok",B{payout_row},"")')
    labels = {**_FUNDING_ROWS, **_EXIT_ROWS, "exit_date": 13, "holding_period_years": 14,
              "status": 16, "failure_period": 17, "funding_shortfall": 18}
    for key, row in labels.items():
        sheet.cell(row, 1, key)
    for row, expression in expressions.items():
        _formula(sheet, f"B{row}", expression)
    for column, label in enumerate(("Cash-flow event", "Actual date", "Deal", "Sponsor", "Management"), 1):
        sheet.cell(43, column, label)
    sheet["A44"], sheet["A45"] = "Entry investment", "Exit distribution"
    _formula(sheet, "B44", _input("deal_date"))
    _formula(sheet, "B45", "B13")
    for address in ("B13", "B44", "B45"):
        sheet[address].number_format = "yyyy-mm-dd"
    for row in (37, 38, 39):
        sheet[f"B{row}"].number_format = _PERCENT_FORMAT
    sheet["A48"] = "Blank valuation or return cells map to JSON null, never fabricated zero returns."
    sheet["A49"] = "Priority: sponsor preferred first, then ordinary split; distributions sum to equity proceeds."
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    workbook.close()


def read_workbook_outputs(path: Path) -> dict:
    """Extract exactly 55 fields from Excel's saved caches; never calculate them."""
    workbook = load_workbook(Path(path), data_only=True, read_only=False, keep_links=False)
    try:
        for name in ("Inputs", "Scenarios", "Forecast", "Debt", "Returns"):
            if name not in workbook.sheetnames:
                raise ValueError(f"Not a Pierre LBO One derivative: missing {name} sheet")
        returns = workbook["Returns"]
        status = returns["B16"].value
        if status not in ("ok", "liquidity_shortfall"):
            raise ValueError("Missing or invalid Excel formula caches; recalculate and save in Microsoft Excel before extraction")

        def number(sheet: str, address: str, nullable: bool = False):
            value = workbook[sheet][address].value
            if nullable and value in (None, ""):
                return None
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                raise ValueError(f"Missing, nonnumeric or nonfinite Excel cache at {sheet}!{address}: {value!r}")
            return float(value)

        def iso(sheet: str, address: str) -> str:
            value = workbook[sheet][address].value
            if isinstance(value, datetime):
                return value.date().isoformat()
            if isinstance(value, date):
                return value.isoformat()
            raise ValueError(f"Missing or invalid Excel date cache at {sheet}!{address}: {value!r}")

        horizon_value = number("Inputs", f"B{_INPUT_ROWS['exit_period']}")
        if not horizon_value.is_integer() or not 1 <= horizon_value <= 8:
            raise ValueError("Invalid exit_period in derivative workbook")
        horizon = int(horizon_value)
        failure = None
        if status == "liquidity_shortfall":
            failure_value = number("Returns", "B17")
            if not failure_value.is_integer() or not 1 <= failure_value <= horizon:
                raise ValueError("Invalid failure_period in derivative workbook")
            failure = int(failure_value)
        completed = horizon if failure is None else failure - 1
        result = {
            "status": status,
            "failure_period": failure,
            "funding_shortfall": number("Returns", "B18"),
            "dates": [iso("Forecast", f"C{row}") for row in range(5, horizon + 5)],
            "operating_fractions": [number("Forecast", f"E{row}") for row in range(5, horizon + 5)],
            "interest_year_fractions": [number("Forecast", f"F{row}") for row in range(5, horizon + 5)],
            "exit_date": iso("Returns", "B13"),
            "holding_period_years": number("Returns", "B14"),
        }
        for key, row in _FUNDING_ROWS.items():
            result[key] = number("Returns", f"B{row}")
        for key, column in _FORECAST_COLUMNS.items():
            result[key] = [number("Forecast", f"{column}{row}") for row in range(5, horizon + 5)]
        for key, column in _FINANCING_COLUMNS.items():
            result[key] = [number("Debt", f"{column}{row}") for row in range(5, completed + 5)]
        for key, row in _EXIT_ROWS.items():
            result[key] = (None if failure is not None else
                           number("Returns", f"B{row}", nullable=key.endswith(("_mom", "_irr"))))
        return result
    finally:
        workbook.close()


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    commands = parser.add_subparsers(dest="command", required=True)
    build = commands.add_parser("build", help="Write an uncalculated, formula-only Excel derivative")
    build.add_argument("--inputs", type=Path, required=True, help="Contract input JSON file")
    build.add_argument("--output", type=Path, required=True, help="Derivative .xlsx destination, never the source")
    extract = commands.add_parser("extract", help="Print JSON from a derivative recalculated and saved by Excel")
    extract.add_argument("--workbook", type=Path, required=True)
    args = parser.parse_args(argv)
    if args.command == "build":
        build_workbook(json.loads(args.inputs.read_text(encoding="utf-8")), args.output)
    else:
        print(json.dumps(read_workbook_outputs(args.workbook), indent=2, allow_nan=False))


if __name__ == "__main__":
    main()
